from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\ASUS\CShape")
OUT_DIR = ROOT / "report_work"
IMG_DIR = OUT_DIR / "figures"
DOCX_PATH = ROOT / "BaoCao_VinhKhanh_CSharp.docx"
IMG_DIR.mkdir(parents=True, exist_ok=True)

PROJECT_TITLE = "Ứng dụng Thuyết minh Đa ngôn ngữ Phố Ẩm thực Vĩnh Khánh"
COURSE = "Môn: Lập trình C#"
UNIVERSITY = "TRƯỜNG ĐẠI HỌC SÀI GÒN"
FACULTY = "KHOA CÔNG NGHỆ THÔNG TIN"

NAVY = RGBColor(31, 58, 95)
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(30, 30, 30)
MUTED = RGBColor(95, 95, 95)
LIGHT_FILL = "E8EEF5"
SOFT_FILL = "F4F6F9"
HEADER_FILL = "D9EAF7"
BORDER = "8EA9C1"


def font_path(name: str) -> str:
    path = Path(r"C:\Windows\Fonts") / name
    return str(path if path.exists() else Path(r"C:\Windows\Fonts\arial.ttf"))


FONT_REG = font_path("arial.ttf")
FONT_BOLD = font_path("arialbd.ttf")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6") -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, width_cm: float, col_widths_cm: list[float] | None = None) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(width_cm / 2.54 * 1440)))
    if col_widths_cm:
        for row in table.rows:
            for idx, width in enumerate(col_widths_cm):
                if idx < len(row.cells):
                    row.cells[idx].width = Cm(width)


def set_run_font(run, size=None, bold=None, color=None, italic=None, name="Arial") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def setup_styles(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2)
    sec.right_margin = Cm(2)
    sec.header_distance = Cm(1.1)
    sec.footer_distance = Cm(1.1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, before, after in [
        ("Title", 18, NAVY, 0, 8),
        ("Heading 1", 15, BLUE, 14, 7),
        ("Heading 2", 12.5, BLUE, 9, 4),
        ("Heading 3", 11.3, NAVY, 6, 3),
    ]:
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = False

    for list_style in ("List Bullet", "List Number"):
        st = styles[list_style]
        st.font.name = "Arial"
        st.font.size = Pt(10)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.1


def set_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = header.add_run("Báo cáo đồ án C# | Vĩnh Khánh Food Tour")
        set_run_font(r, size=8.5, color=MUTED)
        footer = section.footer.paragraphs[0]
        add_page_number(footer)


def add_para(doc, text="", style=None, bold=False, italic=False, align=None, size=None, color=None, after=None) -> None:
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, italic=italic, color=color)


def add_heading(doc, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_caption(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=9, italic=True, color=MUTED)


def add_table(doc, headers: list[str], rows: list[list[str]], widths_cm: list[float] | None = None, font_size=8.8) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_borders(table)
    if widths_cm:
        set_table_width(table, sum(widths_cm), widths_cm)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        set_cell_shading(hdr[i], HEADER_FILL)
        set_cell_margins(hdr[i])
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=font_size, bold=True, color=NAVY)
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            set_cell_margins(cells[i])
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=DARK)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def note_box(doc, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table, color="D0D7DE", size="4")
    set_table_width(table, 16.5, [16.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, SOFT_FILL)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_run_font(r, bold=True, size=10, color=NAVY)
    p.add_run("\n")
    r2 = p.add_run(body)
    set_run_font(r2, size=9.5, color=DARK)


def rounded_rect(draw, box, fill, outline, width=2, radius=18):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def wrap_text(draw, text, font, max_width):
    words = text.split()
    lines, cur = [], ""
    for w in words:
        trial = (cur + " " + w).strip()
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def draw_architecture(path: Path) -> None:
    w, h = 1600, 980
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title_f = ImageFont.truetype(FONT_BOLD, 42)
    box_f = ImageFont.truetype(FONT_BOLD, 28)
    text_f = ImageFont.truetype(FONT_REG, 22)
    small_f = ImageFont.truetype(FONT_REG, 19)
    d.text((50, 35), "Kiến trúc tổng quan hệ thống", fill=(31, 58, 95), font=title_f)

    boxes = [
        ((70, 140, 450, 360), "PWA du khách", ["index.html", "app.js", "map.js / geofence.js", "audio-manager.js"]),
        ((610, 140, 990, 360), "CMS Admin", ["admin.html", "admin.js", "Dashboard", "POI/Tour/Online Users"]),
        ((1150, 140, 1530, 360), "Thiết bị / QR", ["Camera QR", "GPS", "Web Speech", "Service Worker"]),
        ((190, 510, 680, 760), "ASP.NET Core 10", ["Program.cs", "Controllers", "Minimal API demo mode", "/api/tts, /api/presence"]),
        ((900, 510, 1390, 760), "Service & Data Layer", ["AuthService, TourService", "AnalyticsService", "UserPresenceService", "MongoDB Atlas hoặc in-memory demo"]),
    ]
    for box, title, lines in boxes:
        rounded_rect(d, box, fill=(244, 248, 252), outline=(142, 169, 193), width=3)
        d.text((box[0] + 24, box[1] + 22), title, fill=(31, 58, 95), font=box_f)
        y = box[1] + 76
        for line in lines:
            d.text((box[0] + 28, y), "• " + line, fill=(30, 30, 30), font=text_f)
            y += 36

    arrows = [
        ((450, 245), (610, 245), "REST API"),
        ((990, 245), (1150, 245), "QR / URL"),
        ((330, 360), (430, 510), "fetch"),
        ((800, 360), (610, 510), "admin API"),
        ((680, 635), (900, 635), "DI Services"),
    ]
    for (x1, y1), (x2, y2), label in arrows:
        d.line((x1, y1, x2, y2), fill=(46, 116, 181), width=5)
        d.polygon([(x2, y2), (x2 - 18, y2 - 10), (x2 - 18, y2 + 10)], fill=(46, 116, 181))
        mx, my = (x1 + x2) // 2, (y1 + y2) // 2
        d.text((mx - 55, my - 32), label, fill=(46, 116, 181), font=small_f)

    d.text((70, 860), "Ghi chú: Khi chưa cấu hình MongoDB, backend chuyển sang demo in-memory để phục vụ bảo vệ đồ án.", fill=(95, 95, 95), font=text_f)
    img.save(path)


def draw_sequence(path: Path, title: str, participants: list[str], messages: list[tuple[int, int, str]]) -> None:
    width = 1700
    top = 120
    lane_gap = width // (len(participants) + 1)
    height = max(740, top + len(messages) * 70 + 130)
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    title_f = ImageFont.truetype(FONT_BOLD, 38)
    part_f = ImageFont.truetype(FONT_BOLD, 22)
    msg_f = ImageFont.truetype(FONT_REG, 19)
    d.text((50, 36), title, fill=(31, 58, 95), font=title_f)

    xs = []
    for idx, p in enumerate(participants):
        x = lane_gap * (idx + 1)
        xs.append(x)
        rounded_rect(d, (x - 120, top, x + 120, top + 58), fill=(232, 238, 245), outline=(142, 169, 193), radius=14)
        lines = wrap_text(d, p, part_f, 210)
        yy = top + 9
        for line in lines[:2]:
            tw = d.textbbox((0, 0), line, font=part_f)[2]
            d.text((x - tw / 2, yy), line, fill=(31, 58, 95), font=part_f)
            yy += 24
        d.line((x, top + 58, x, height - 60), fill=(170, 185, 200), width=2)

    y = top + 100
    for src, dst, msg in messages:
        x1, x2 = xs[src], xs[dst]
        d.line((x1, y, x2, y), fill=(46, 116, 181), width=3)
        direction = 1 if x2 > x1 else -1
        d.polygon([(x2, y), (x2 - direction * 16, y - 8), (x2 - direction * 16, y + 8)], fill=(46, 116, 181))
        label_lines = wrap_text(d, msg, msg_f, abs(x2 - x1) - 30 if abs(x2 - x1) > 240 else 300)
        label = label_lines[0] if label_lines else msg
        d.text((min(x1, x2) + 18, y - 28), label, fill=(30, 30, 30), font=msg_f)
        y += 70

    img.save(path)


def draw_activity(path: Path) -> None:
    w, h = 1500, 960
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    title_f = ImageFont.truetype(FONT_BOLD, 38)
    step_f = ImageFont.truetype(FONT_BOLD, 24)
    text_f = ImageFont.truetype(FONT_REG, 20)
    d.text((50, 35), "Luồng khởi động demo mode và MongoDB mode", fill=(31, 58, 95), font=title_f)
    steps = [
        ("Start", "Program.cs đọc cấu hình"),
        ("Kiểm tra MongoDB", "Có connection string hợp lệ?"),
        ("MongoDB mode", "MapControllers + Services dùng MongoDB"),
        ("Demo mode", "MapDemoApi + demoPois/demoTours/analyticsEvents"),
        ("Frontend", "index.html/admin.html gọi API theo cùng contract"),
        ("Ready", "App sẵn sàng demo QR, dashboard, audio, online users"),
    ]
    coords = [(120, 140), (520, 140), (200, 410), (840, 410), (520, 650), (520, 830)]
    for (label, desc), (x, y) in zip(steps, coords):
        rounded_rect(d, (x, y, x + 330, y + 110), fill=(244, 248, 252), outline=(142, 169, 193), radius=18)
        d.text((x + 20, y + 18), label, fill=(31, 58, 95), font=step_f)
        for i, line in enumerate(wrap_text(d, desc, text_f, 285)[:2]):
            d.text((x + 20, y + 55 + i * 24), line, fill=(30, 30, 30), font=text_f)
    for a, b in [(0, 1), (1, 2), (1, 3), (2, 4), (3, 4), (4, 5)]:
        x1, y1 = coords[a][0] + 330, coords[a][1] + 55
        x2, y2 = coords[b][0], coords[b][1] + 55
        if b in (2, 3):
            x1, y1 = coords[a][0] + 165, coords[a][1] + 110
            x2, y2 = coords[b][0] + 165, coords[b][1]
        if a in (2, 3):
            x1, y1 = coords[a][0] + 165, coords[a][1] + 110
            x2, y2 = coords[b][0] + 165, coords[b][1]
        d.line((x1, y1, x2, y2), fill=(46, 116, 181), width=4)
        d.polygon([(x2, y2), (x2 - 10, y2 - 18), (x2 + 10, y2 - 18)], fill=(46, 116, 181))
    img.save(path)


def create_diagrams() -> dict[str, Path]:
    paths = {
        "architecture": IMG_DIR / "architecture.png",
        "startup": IMG_DIR / "startup.png",
        "qr": IMG_DIR / "qr_sequence.png",
        "dashboard": IMG_DIR / "dashboard_sequence.png",
        "kick": IMG_DIR / "kick_sequence.png",
        "audio": IMG_DIR / "audio_sequence.png",
    }
    draw_architecture(paths["architecture"])
    draw_activity(paths["startup"])
    draw_sequence(
        paths["qr"],
        "Sequence QR Tour tại cổng",
        ["Du khách", "QRScanner", "app.js", "API Tour", "TourService", "Analytics"],
        [
            (0, 1, "Quét QR Tour"),
            (1, 2, "QRScannerManager._onScanSuccess()"),
            (2, 2, "handleQrCode() -> handleTourQrCode()"),
            (2, 2, "extractTourQrCode()"),
            (2, 3, "resolveTourByQrCode(): GET /api/tour/qr/{qrCode}"),
            (3, 4, "GetActiveByQrCodeAsync()"),
            (4, 4, "GetOrderedPoisAsync()"),
            (3, 2, "JSON tour + pois[]"),
            (2, 2, "openTourFromQr() -> renderTourPoiList()"),
            (2, 2, "showPoiDetail()"),
            (2, 5, "trackTourQrScan(): qr_scan"),
        ],
    )
    draw_sequence(
        paths["dashboard"],
        "Sequence Dashboard Analytics",
        ["Admin", "admin.js", "Stats API", "Top/Recent API", "Heatmap API"],
        [
            (0, 1, "Mở Dashboard"),
            (1, 2, "loadDashboardData(): GET /api/analytics/stats"),
            (1, 3, "GET /api/analytics/top-pois"),
            (1, 3, "GET /api/analytics/recent"),
            (1, 4, "initHeatmap(): GET /api/analytics/heatmap"),
            (2, 1, "eventCounts + uniqueSessions"),
            (3, 1, "topPois[] + recentEvents[]"),
            (4, 1, "points[]"),
            (1, 1, "renderTopPoisChart()"),
            (1, 1, "renderRecentEvents()"),
        ],
    )
    draw_sequence(
        paths["kick"],
        "Sequence Online Users và Kick Session",
        ["Client app", "Presence API", "UserPresenceService", "Admin CMS", "admin.js"],
        [
            (0, 0, "startPresenceHeartbeat()"),
            (0, 1, "sendPresenceHeartbeat(): POST /api/presence/heartbeat"),
            (1, 2, "Upsert()"),
            (3, 4, "loadOnlineUsers()"),
            (4, 1, "GET /api/admin/online-users"),
            (1, 2, "GetOnlineSessions()"),
            (3, 4, "kickOnlineUser()"),
            (4, 1, "POST /api/admin/online-users/{sessionId}/kick"),
            (1, 2, "Kick()"),
            (0, 1, "Heartbeat kế tiếp"),
            (1, 0, "kicked = true"),
            (0, 0, "handleSessionKicked()"),
        ],
    )
    draw_sequence(
        paths["audio"],
        "Sequence Audio, TTS và đa ngôn ngữ",
        ["Du khách", "app.js", "Translation", "AudioManager", "/api/tts"],
        [
            (0, 1, "Đổi ngôn ngữ"),
            (1, 1, "changeLanguage() -> normalizeAppLanguage()"),
            (1, 1, "applyUILanguage()"),
            (1, 2, "translateWithCache() -> translateText()"),
            (1, 1, "queuePoiTranslationWarmup()"),
            (0, 1, "Bấm nghe"),
            (1, 1, "getPoiScript()"),
            (1, 3, "AudioManager.playDirect()"),
            (3, 3, "_speak()"),
            (3, 3, "_speakWithWebSpeech()"),
            (3, 4, "Fallback _speakWithGoogleTTS(): GET /api/tts"),
        ],
    )
    return paths


def add_cover(doc: Document) -> None:
    for _ in range(2):
        add_para(doc, "")
    add_para(doc, UNIVERSITY, bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, FACULTY, bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "", after=28)
    add_para(doc, COURSE, bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "BÁO CÁO ĐỒ ÁN CUỐI KỲ", bold=True, size=17, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Đề tài:", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, PROJECT_TITLE, bold=True, size=17, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
    add_para(doc, "", after=26)
    add_para(doc, "Thành viên nhóm: [Điền tên thành viên]", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Mã số sinh viên: [Điền mã số sinh viên]", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Giảng viên phụ trách: [Điền tên giảng viên]", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "", after=84)
    add_para(doc, "TP. Hồ Chí Minh, năm 2026", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    add_heading(doc, "MỤC LỤC TÓM TẮT", 1)
    chapters = [
        "CHƯƠNG 1: Kiến trúc tổng quan",
        "CHƯƠNG 2: Luồng khởi động hệ thống",
        "CHƯƠNG 3: QR Tour và tương tác O2O",
        "CHƯƠNG 4: Bản đồ, GPS và Geofence",
        "CHƯƠNG 5: Audio, TTS và đa ngôn ngữ",
        "CHƯƠNG 6: Admin Authentication và quản trị",
        "CHƯƠNG 7: Dashboard Analytics",
        "CHƯƠNG 8: Online Users và Kick Session",
        "CHƯƠNG 9: Offline/PWA và dữ liệu demo",
        "CHƯƠNG 10: End-to-End Flow",
        "CHƯƠNG 11: Bảng đặc tả hệ sinh thái chức năng",
        "CHƯƠNG 12: Design Patterns & Key Constants",
        "PHỤ LỤC: Công nghệ, chi phí, kiểm thử và gợi ý bảo vệ",
    ]
    for item in chapters:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)
    doc.add_page_break()


def chapter_1(doc, figs):
    add_heading(doc, "CHƯƠNG 1: KIẾN TRÚC TỔNG QUAN", 1)
    add_para(doc, "Hệ thống được xây dựng theo mô hình web app phục vụ hai nhóm người dùng: du khách sử dụng PWA để quét QR, xem bản đồ và nghe thuyết minh; quản trị viên sử dụng CMS để quản lý nội dung, theo dõi analytics và điều khiển phiên online. Backend ASP.NET Core 10 đóng vai trò API trung tâm, đồng thời hỗ trợ hai chế độ dữ liệu: MongoDB Atlas và demo in-memory.")
    doc.add_picture(str(figs["architecture"]), width=Cm(16.5))
    add_caption(doc, "Hình 1.1. Kiến trúc tổng quan PWA - CMS - ASP.NET Core - MongoDB/demo mode")
    add_heading(doc, "1.1. Các tầng chính", 2)
    add_table(doc, ["Tầng", "Thành phần", "Trách nhiệm"], [
        ["Frontend du khách", "index.html, app.js, map.js, geofence.js, audio-manager.js", "Hiển thị bản đồ, QR Tour, POI detail, đổi ngôn ngữ, phát audio, cache offline."],
        ["CMS Admin", "admin.html, admin.js", "Đăng nhập admin, CRUD POI/Tour, tạo QR, dashboard, online users, kick session."],
        ["Backend API", "Program.cs, Controllers/*", "Bootstrap hệ thống, route API, minimal API demo mode, controller MongoDB mode."],
        ["Service layer", "AuthService, TourService, AnalyticsService, UserPresenceService", "Xử lý nghiệp vụ đăng nhập, resolve tour, thống kê, heartbeat/kick."],
        ["Data layer", "MongoDB Atlas hoặc biến in-memory", "Lưu pois, tours, analytics, users; demo mode dùng demoPois, demoTours, analyticsEvents."],
    ], [3.0, 5.6, 7.9], 8.5)
    add_heading(doc, "1.2. Các file quan trọng", 2)
    add_table(doc, ["File", "Vai trò"], [
        ["Program.cs", "Cấu hình pipeline, DI, route demo, /api/tts, /api/presence, /api/admin/online-users, seed demo analytics."],
        ["wwwroot/js/app.js", "Điều phối app du khách: QR, map, i18n, presence, audio, offline, AAC."],
        ["wwwroot/js/admin.js", "Điều phối CMS: login, dashboard, CRUD, QR viewer, online users."],
        ["Services/AuthService.cs", "Hash mật khẩu, xác thực admin, seed user admin bằng hash."],
        ["Services/TourService.cs", "Resolve QR Tour và sắp xếp POI theo thứ tự tour."],
        ["Services/AnalyticsService.cs", "Ghi event, tính top POI, heatmap, recent, unique sessions."],
        ["Services/UserPresenceService.cs", "Lưu heartbeat in-memory, xác định online window 45 giây và trạng thái kick."],
    ], [5.2, 11.3], 9)
    note_box(doc, "Phạm vi bảo mật", "Báo cáo chỉ trình bày bảo mật mức demo: mật khẩu admin không lưu plaintext mà được băm trước khi so sánh. Hệ thống chưa thiết kế refresh token, rate limit, CSRF hoặc phân quyền phức tạp.")


def chapter_2(doc, figs):
    add_heading(doc, "CHƯƠNG 2: LUỒNG KHỞI ĐỘNG HỆ THỐNG", 1)
    add_para(doc, "Khi backend khởi động, Program.cs đọc cấu hình, kiểm tra MongoDB__ConnectionString hoặc appsettings.Local.json, sau đó quyết định chạy MongoDB mode hay demo mode. Cách này giúp đồ án vẫn demo được khi môi trường không có MongoDB, đồng thời giữ cùng contract API cho frontend.")
    doc.add_picture(str(figs["startup"]), width=Cm(16.0))
    add_caption(doc, "Hình 2.1. Luồng quyết định MongoDB mode và demo mode")
    add_heading(doc, "2.1. Backend startup", 2)
    add_numbered(doc, [
        "Đọc cấu hình URL, MongoDB connection string và database name.",
        "Đăng ký controller, OpenAPI, CORS, static files và service cần thiết.",
        "Nếu có MongoDB, đăng ký IMongoClient, IMongoDatabase và các service CRUD/analytics/auth.",
        "Nếu không có MongoDB, gọi MapDemoApi(app) để đăng ký dữ liệu in-memory.",
        "Map route gốc / sang /index.html, sau đó phục vụ app du khách và CMS admin.",
    ])
    add_heading(doc, "2.2. Frontend startup", 2)
    add_table(doc, ["Luồng", "Hàm/tệp", "Mô tả"], [
        ["User app", "index.html -> app.js", "Khởi tạo AppState, ngôn ngữ, map, geofence, audioManager, offlineDB và heartbeat."],
        ["Admin CMS", "admin.html -> admin.js", "Kiểm tra session admin, gắn listener login, nạp dashboard hoặc trang đang chọn."],
        ["Offline", "offline-db.js, sw.js", "IndexedDB lưu POI; service worker cache app shell, CDN, tiles và API GET."],
    ], [3.3, 5.5, 7.7], 8.7)
    add_heading(doc, "2.3. Demo mode vs MongoDB mode", 2)
    add_para(doc, "Điểm quan trọng của kiến trúc là frontend không phải đổi luồng khi backend chuyển mode. Ví dụ /api/tour/qr/{qrCode}, /api/analytics/stats, /api/poi/all vẫn giữ endpoint giống nhau. Khác biệt nằm ở nguồn dữ liệu: MongoDB collection hoặc danh sách in-memory được seed trong Program.cs.")


def chapter_3(doc, figs):
    add_heading(doc, "CHƯƠNG 3: QR TOUR VÀ TƯƠNG TÁC O2O", 1)
    add_para(doc, "Luồng O2O của đồ án bắt đầu từ QR Tour đặt tại cổng hoặc trên CMS. QR không chỉ mở một POI đơn lẻ mà mở toàn bộ tour, gồm danh sách quán theo thứ tự được cấu hình trong tours.poiIds. Cách này phù hợp tình huống thực tế: du khách đến cổng phố ẩm thực, quét mã và đi theo tuyến gợi ý.")
    add_heading(doc, "3.1. Luồng QR Tour", 2)
    add_table(doc, ["Bước", "Hàm/API", "Ý nghĩa"], [
        ["1", "QRScannerManager._onScanSuccess()", "Scanner nhận nội dung QR từ camera."],
        ["2", "handleQrCode() -> handleTourQrCode()", "Phân nhánh QR Tour trước QR POI cũ."],
        ["3", "extractTourQrCode()", "Rút qrCode từ URL hoặc chuỗi raw."],
        ["4", "resolveTourByQrCode()", "Gọi GET /api/tour/qr/{qrCode}."],
        ["5", "openTourFromQr()", "Gán activeTour và nạp danh sách POI."],
        ["6", "renderTourPoiList() + showPoiDetail()", "Hiển thị tour list và mở điểm đầu tiên."],
        ["7", "trackTourQrScan()", "Ghi analytics eventType = qr_scan."],
    ], [1.2, 6.2, 9.1], 8.7)
    add_heading(doc, "3.2. Ordered tour list", 2)
    add_para(doc, "TourService.GetOrderedPoisAsync() dùng thứ tự trong tour.poiIds để sắp xếp POI. Trong demo mode, Program.cs cũng tạo dictionary id -> index để trả về pois theo đúng thứ tự. Đây là chi tiết then chốt để khi quét QR tại cổng, danh sách quán hiển thị theo tuyến đã thiết kế, không phụ thuộc thứ tự lưu trong database.")
    doc.add_picture(str(figs["qr"]), width=Cm(14.2))
    add_caption(doc, "Hình 3.1. Sequence QR Tour dùng hàm thật trong dự án")


def chapter_4(doc):
    add_heading(doc, "CHƯƠNG 4: BẢN ĐỒ, GPS VÀ GEOFENCE", 1)
    add_para(doc, "App du khách hiển thị bản đồ bằng Leaflet/OpenStreetMap và đặt marker cho các POI. Khi người dùng bấm marker hoặc chọn từ danh sách tour, app gọi showPoiDetail() để mở nội dung thuyết minh, địa chỉ, giờ mở cửa, khoảng giá và nút nghe audio.")
    add_heading(doc, "4.1. Map rendering và POI detail", 2)
    add_bullets(doc, [
        "Dữ liệu POI được lấy từ GET /api/poi khi online, hoặc từ IndexedDB/demo fallback khi API lỗi.",
        "Marker theo category giúp phân biệt hải sản, lẩu, street food, snack và landmark.",
        "showPoiDetail() là điểm hội tụ của nhiều luồng: click marker, chọn trong tour list, quét QR và điều hướng Next/Previous.",
    ])
    add_heading(doc, "4.2. GPS/geofence trigger", 2)
    add_para(doc, "Geofence được dùng để gợi ý nghe khi du khách đến gần POI. Vì GPS ở phố nhỏ có thể lệch, app không phụ thuộc tuyệt đối vào GPS; QR Tour và QR điểm được dùng làm fallback thực tế hơn khi bảo vệ/demo.")
    add_table(doc, ["Tình huống", "Cách xử lý"], [
        ["GPS hoạt động", "Bắt vị trí, tính khoảng cách tới POI, hiển thị gợi ý nếu trong bán kính."],
        ["GPS bị chặn", "App vẫn cho mở tour bằng QR, chọn POI từ list, nghe thuyết minh thủ công."],
        ["API lỗi", "Dùng offlineDB.loadPois() hoặc getDemoPois() để tránh trắng màn hình."],
    ], [4.2, 12.3], 9)


def chapter_5(doc, figs):
    add_heading(doc, "CHƯƠNG 5: AUDIO, TTS VÀ ĐA NGÔN NGỮ", 1)
    add_para(doc, "Hệ thống hỗ trợ 20 ngôn ngữ theo mã vi, en, ja, zh, ko, th, fr, es, de, ru, pt, it, id, hi, ar, ms, tl, nl, sv, pl. App ưu tiên dữ liệu đã có trong POI; nếu thiếu bản dịch, luồng translateWithCache() gọi translateText() và cache lại để giảm gọi lặp.")
    add_heading(doc, "5.1. Language switching", 2)
    add_table(doc, ["Hàm", "Vai trò"], [
        ["changeLanguage()", "Điểm vào khi người dùng đổi dropdown ngôn ngữ."],
        ["normalizeAppLanguage()", "Chuẩn hóa mã ngôn ngữ, ví dụ jp -> ja nếu cần."],
        ["applyUILanguage()", "Đổi text UI như nút nghe, đóng, chỉ đường, tour step."],
        ["translateWithCache()", "Cache bản dịch theo scope/id/source/target."],
        ["queuePoiTranslationWarmup()", "Làm nóng dịch POI để thao tác sau nhanh hơn."],
    ], [5.2, 11.3], 9)
    add_heading(doc, "5.2. Audio fallback", 2)
    add_para(doc, "AudioManager.playDirect() gọi _speak(). Nếu browser có voice phù hợp, app dùng _speakWithWebSpeech(). Nếu thiếu voice hoặc phát sinh lỗi, app fallback sang _speakWithGoogleTTS() và endpoint /api/tts. Đây là thiết kế phù hợp cho demo trên nhiều máy vì Web Speech phụ thuộc trình duyệt và hệ điều hành.")
    doc.add_picture(str(figs["audio"]), width=Cm(16.3))
    add_caption(doc, "Hình 5.1. Sequence đổi ngôn ngữ, lấy script và phát audio")


def chapter_6(doc):
    add_heading(doc, "CHƯƠNG 6: ADMIN AUTHENTICATION VÀ QUẢN TRỊ", 1)
    add_para(doc, "CMS Admin gồm form đăng nhập, dashboard và các màn quản trị POI/Tour. Mật khẩu admin không được lưu plaintext trong seed; AuthService.ComputePasswordHash() băm password đầu vào và VerifyPassword() so sánh với hash đã lưu.")
    add_heading(doc, "6.1. Login flow", 2)
    add_table(doc, ["Bước", "Thành phần", "Mô tả"], [
        ["1", "admin.html", "Admin nhập username/password và submit form."],
        ["2", "loginAdmin()", "Gửi POST /api/auth/login."],
        ["3", "AuthController.Login() hoặc demo route Program.cs", "Nhận LoginRequest và chuyển sang AuthService hoặc VerifyPassword."],
        ["4", "AuthService.LoginAsync()", "Tính hash password, tìm user admin, kiểm tra role/isActive."],
        ["5", "AdminTokenHelper", "Trả token admin cho CMS lưu trong sessionStorage."],
    ], [1.0, 5.1, 10.4], 8.8)
    note_box(doc, "Ghi chú bảo mật", "Đây là bảo mật mức demo phục vụ đồ án: có hash mật khẩu và token admin đơn giản. Báo cáo không khẳng định hệ thống đã đạt chuẩn production về refresh token, rate limit, CSRF hay phân quyền nhiều cấp.")
    add_heading(doc, "6.2. CRUD POI/Tour", 2)
    add_para(doc, "POI và Tour có đủ luồng đọc, tạo, sửa, xóa ở MongoDB mode qua controller/service; demo mode cũng map POST/PUT/DELETE trong Program.cs. Frontend dùng savePoi() và saveTour(), đồng thời giữ QR code để in hoặc quét lại.")


def chapter_7(doc, figs):
    add_heading(doc, "CHƯƠNG 7: DASHBOARD ANALYTICS", 1)
    add_para(doc, "Dashboard CMS tổng hợp các sự kiện chính: qr_scan, poi_listen, location_update và số unique sessions. Các API được load trong loadDashboardData(), sau đó renderTopPoisChart(), renderRecentEvents() và initHeatmap() cập nhật giao diện.")
    add_table(doc, ["Chỉ số", "Nguồn", "Cách tính/hiển thị"], [
        ["QR Scans", "eventCounts.qr_scan", "Tăng khi trackTourQrScan() ghi event quét tour."],
        ["Lượt nghe", "eventCounts.poi_listen", "Tăng khi audio manager ghi nhận nghe POI."],
        ["Sessions", "GetUniqueSessionsAsync()", "Đếm sessionId khác nhau trong analytics."],
        ["Top POIs", "GetTopPoiStatsAsync()", "Group event poi_listen theo poiId và sắp theo ListenCount."],
        ["Heatmap", "GetHeatmapDataAsync()", "Dùng event vị trí để tạo điểm heatmap."],
    ], [3.2, 5.8, 7.5], 8.8)
    doc.add_picture(str(figs["dashboard"]), width=Cm(16.3))
    add_caption(doc, "Hình 7.1. Sequence Dashboard Analytics")
    add_heading(doc, "7.1. Demo data", 2)
    add_para(doc, "Trong demo mode, CreateDemoAnalyticsEvents() seed sẵn một số event để dashboard không trắng khi chưa có MongoDB. Khi demo thật, việc quét QR hoặc bấm nghe sẽ tiếp tục ghi event mới vào analyticsEvents in-memory.")


def chapter_8(doc, figs):
    add_heading(doc, "CHƯƠNG 8: ONLINE USERS VÀ KICK SESSION", 1)
    add_para(doc, "Online Users là chức năng demo quản trị phiên đang truy cập app du khách. Client gửi heartbeat mỗi 15 giây; backend lưu session trong UserPresenceService. Một user được tính online nếu LastSeenAt nằm trong cửa sổ 45 giây và chưa bị kick.")
    add_table(doc, ["Hằng số/logic", "Giá trị", "Ý nghĩa"], [
        ["OnlineWindow", "45 giây", "Heartbeat trong khoảng này được xem là online."],
        ["RetentionWindow", "30 phút", "Session cũ được giữ trong RAM trước khi cleanup."],
        ["UserPresenceService", "Singleton", "Dữ liệu online dùng chung trong một instance backend."],
        ["Kick()", "IsKicked = true", "Admin đánh dấu session bị ngắt."],
        ["handleSessionKicked()", "Overlay client", "Client dừng audio/GPS và hiện màn bị ngắt."],
    ], [4.4, 3.1, 9.0], 8.8)
    doc.add_picture(str(figs["kick"]), width=Cm(16.3))
    add_caption(doc, "Hình 8.1. Sequence Online Users và Kick Session")
    note_box(doc, "Giới hạn user online", "Code hiện tại không đặt giới hạn cứng số user online. Giới hạn thực tế phụ thuộc RAM, CPU, network và cách deploy. Với bản demo, online users là dữ liệu in-memory, phù hợp trình bày logic hơn là tải production.")


def chapter_9(doc):
    add_heading(doc, "CHƯƠNG 9: OFFLINE/PWA VÀ DỮ LIỆU DEMO", 1)
    add_para(doc, "App được thiết kế như PWA: service worker cache app shell, CDN cần thiết và API GET; IndexedDB lưu POI để app vẫn có dữ liệu khi mất mạng. Đây là lớp dự phòng quan trọng cho bối cảnh demo bằng điện thoại và mạng WiFi không ổn định.")
    add_heading(doc, "9.1. Service Worker", 2)
    add_table(doc, ["Nhóm tài nguyên", "Chiến lược", "Ví dụ"], [
        ["App shell", "CacheFirst", "index.html, admin.html, css, js, manifest.json."],
        ["API GET", "NetworkFirst", "Ưu tiên dữ liệu mới, lỗi mạng thì dùng cache."],
        ["Map tiles", "CacheFirst", "Tile OpenStreetMap khu vực Quận 4."],
        ["CDN", "CacheFirst", "Leaflet, html5-qrcode, QRCode, fonts."],
    ], [3.5, 4.0, 9.0], 9)
    add_heading(doc, "9.2. Demo seed data", 2)
    add_bullets(doc, [
        "CreateDemoPois() tạo danh sách điểm ăn uống, landmark và điểm dừng tour.",
        "CreateDemoTours() tạo QR Tour cổng với mã VK-TOUR-GATE-001.",
        "CreateDemoAnalyticsEvents() tạo dữ liệu mẫu cho QR scans, lượt nghe và heatmap.",
        "Demo mode giúp bảo vệ đồ án kể cả khi chưa có MongoDB Atlas hoặc mạng không ổn định.",
    ])


def chapter_10(doc, figs):
    add_heading(doc, "CHƯƠNG 10: END-TO-END FLOW", 1)
    add_para(doc, "Chương này gom ba luồng quan trọng nhất khi bảo vệ: QR Tour, Dashboard và Online/Kick. Các sơ đồ dùng đúng tên hàm thật để khi giảng viên hỏi có thể truy từ UI đến API và service.")
    doc.add_picture(str(figs["qr"]), width=Cm(14.8))
    add_caption(doc, "Hình 10.1. End-to-end QR Tour")
    doc.add_picture(str(figs["dashboard"]), width=Cm(14.8))
    add_caption(doc, "Hình 10.2. End-to-end Dashboard Analytics")
    doc.add_picture(str(figs["kick"]), width=Cm(14.8))
    add_caption(doc, "Hình 10.3. End-to-end Online Users/Kick")


def chapter_11(doc):
    add_heading(doc, "CHƯƠNG 11: BẢNG ĐẶC TẢ HỆ SINH THÁI CHỨC NĂNG", 1)
    add_para(doc, "Bảng dưới đây là bản truy vết cô đọng theo yêu cầu báo cáo: từ thao tác người dùng đến hàm frontend, endpoint, service và nguồn dữ liệu.")
    rows = [
        ["Admin login", "Nhập tài khoản", "loginAdmin()", "POST /api/auth/login", "AuthController/LoginAsync hoặc VerifyPassword", "users.passwordHash", "Trình bày hash mật khẩu."],
        ["CRUD POI", "Thêm/sửa/xóa POI", "savePoi()", "POST/PUT/DELETE /api/poi", "PoiController + PoiService", "pois/demoPois", "Kiểm tra bảng POI và QR."],
        ["CRUD Tour", "Tạo/sửa tour", "saveTour()", "POST/PUT/DELETE /api/tour", "TourController + TourService", "tours/demoTours", "Chọn nhiều POI theo thứ tự."],
        ["QR Tour", "Quét QR cổng", "handleTourQrCode()", "GET /api/tour/qr/{qrCode}", "GetActiveByQrCodeAsync + GetOrderedPoisAsync", "tours, pois", "Mở tour 10 điểm."],
        ["Audio", "Bấm nghe", "getPoiScript()", "GET /api/tts", "AudioManager + TTS route", "pois.ttsScript", "Fallback khi thiếu voice."],
        ["Ngôn ngữ", "Đổi dropdown", "changeLanguage()", "Không bắt buộc", "translateText/cache", "localStorage/cache", "20 ngôn ngữ."],
        ["Dashboard", "Mở Dashboard", "loadDashboardData()", "/api/analytics/*", "AnalyticsController/Service", "analytics", "Stats/top/recent/heatmap."],
        ["Online", "Mở app user", "sendPresenceHeartbeat()", "POST /api/presence/heartbeat", "UserPresenceService.Upsert()", "in-memory sessions", "Online window 45 giây."],
        ["Kick", "Admin bấm Kick", "kickOnlineUser()", "POST /api/admin/online-users/{id}/kick", "UserPresenceService.Kick()", "in-memory sessions", "Client hiện màn bị ngắt."],
        ["Offline/PWA", "Mất mạng", "offlineDB.loadPois()", "GET /api/poi khi online", "Service worker + OfflineDB", "IndexedDB/cache", "Không trắng dữ liệu."],
        ["AAC", "Nói giúp tôi", "aacSpeak()", "Google TTS fallback", "detectLanguage()", "Câu mẫu frontend", "Hỗ trợ du khách giao tiếp."],
    ]
    add_table(doc, ["Chức năng", "User action", "Frontend", "Endpoint", "Backend/service", "Data", "Demo note"], rows, [2.4, 2.4, 2.5, 2.5, 2.7, 2.0, 2.0], 7.0)


def chapter_12(doc):
    add_heading(doc, "CHƯƠNG 12: DESIGN PATTERNS & KEY CONSTANTS", 1)
    add_heading(doc, "12.1. Pattern table", 2)
    add_table(doc, ["Pattern", "Áp dụng trong dự án", "Lợi ích"], [
        ["Layered architecture", "Frontend, Controller/API, Service, Data", "Tách UI khỏi nghiệp vụ và dữ liệu."],
        ["Repository/service style", "PoiService, TourService, AnalyticsService", "Gom logic nghiệp vụ quanh collection liên quan."],
        ["Fallback strategy", "Demo mode, offline cache, TTS fallback", "Demo ổn định khi DB/mạng/browser voice lỗi."],
        ["Cache-aside", "translateWithCache(), OfflineDB, service worker", "Giảm gọi lặp và hỗ trợ offline."],
        ["Heartbeat presence", "sendPresenceHeartbeat() + UserPresenceService", "Theo dõi online users đơn giản, dễ demo."],
    ], [4.1, 6.5, 5.9], 8.8)
    add_heading(doc, "12.2. Key constants", 2)
    add_table(doc, ["Hằng số/giá trị", "Nơi dùng", "Ý nghĩa"], [
        ["20 ngôn ngữ", "SUPPORTED_LANGUAGES", "Danh sách ngôn ngữ app hỗ trợ."],
        ["15 giây", "setInterval(sendPresenceHeartbeat, 15000)", "Chu kỳ heartbeat client."],
        ["45 giây", "OnlineWindow", "Cửa sổ xác định online."],
        ["30 phút", "RetentionWindow", "Thời gian giữ session trong RAM."],
        ["VK-TOUR-GATE-001", "demoTours", "QR Tour chính tại cổng."],
        ["CACHE_NAME/API_CACHE/TILE_CACHE", "sw.js", "Phiên bản cache PWA."],
        ["DefaultAdminPasswordHash", "AuthService", "Hash admin mặc định, không lưu plaintext."],
    ], [4.6, 5.9, 6.0], 8.8)


def appendix(doc):
    add_heading(doc, "PHỤ LỤC", 1)
    add_heading(doc, "A. Technology stack", 2)
    add_table(doc, ["Nhóm", "Công nghệ"], [
        ["Backend", "ASP.NET Core 10, C#, Minimal API, MVC Controller, OpenAPI."],
        ["Database", "MongoDB Atlas; demo in-memory khi không có connection string."],
        ["Frontend", "HTML/CSS/Vanilla JavaScript, PWA, Service Worker, IndexedDB."],
        ["Map/QR", "Leaflet/OpenStreetMap, html5-qrcode, qrcode.js."],
        ["Audio", "Web Speech API, Google TTS fallback qua /api/tts."],
        ["Deploy/demo", "dotnet run, HTTPS localhost/LAN cho GPS trên điện thoại."],
    ], [4.2, 12.3], 9)
    add_heading(doc, "B. Cost / free-tier analysis", 2)
    add_para(doc, "Bản demo có thể vận hành với chi phí 0 đồng nếu dùng MongoDB Atlas free tier, trình duyệt hỗ trợ Web Speech, OpenStreetMap/Leaflet và cache cục bộ. Tuy nhiên free tier không đồng nghĩa SLA production. Endpoint TTS fallback phụ thuộc mạng và dịch vụ bên ngoài, nên khi triển khai thật cần thay bằng API có quota và điều khoản rõ ràng.")
    add_heading(doc, "C. Testing commands", 2)
    add_table(doc, ["Lệnh", "Mục tiêu"], [
        ["node --check CShape/VinhKhanhFoodTour.Api/wwwroot/js/app.js", "Kiểm tra cú pháp JavaScript app du khách."],
        ["node --check CShape/VinhKhanhFoodTour.Api/wwwroot/js/admin.js", "Kiểm tra cú pháp JavaScript CMS admin."],
        ["dotnet build CShape/VinhKhanhFoodTour.Api/VinhKhanhFoodTour.Api.csproj", "Biên dịch backend ASP.NET Core."],
        ["dotnet run --project CShape/VinhKhanhFoodTour.Api/VinhKhanhFoodTour.Api.csproj", "Chạy server local để demo."],
    ], [9.8, 6.7], 8.7)
    note_box(doc, "Lỗi build thường gặp", "Nếu dotnet build báo file .exe bị khóa bởi VinhKhanhFoodTour.Api (PID), nghĩa là server đang chạy. Dừng terminal dotnet run hoặc Stop-Process theo PID rồi build lại.")
    add_heading(doc, "D. Demo talking points", 2)
    add_bullets(doc, [
        "Số QR Scans lấy từ eventCounts.qr_scan trong API /api/analytics/stats.",
        "Top POI nghe nhiều nhất là ListenCount, được group từ event poi_listen.",
        "Quét QR gọi handleQrCode() -> handleTourQrCode() -> resolveTourByQrCode() -> GET /api/tour/qr/{qrCode}.",
        "Đổi ngôn ngữ gọi changeLanguage() -> translateWithCache() -> applyUILanguage() và queue warmup.",
        "Kick user: admin gọi API kick; heartbeat sau đó trả kicked = true; client chạy handleSessionKicked().",
        "Password admin không lưu plaintext; hệ thống băm password đầu vào rồi so sánh hash.",
    ])


def chapter_4(doc):
    add_heading(doc, "CHƯƠNG 4: BẢN ĐỒ, GPS VÀ GEOFENCE", 1)
    add_para(doc, "Chức năng bản đồ là lớp tương tác trực quan của app du khách. Bản đồ không chỉ hiển thị vị trí quán ăn mà còn nối các luồng QR Tour, POI detail, audio và analytics. Trong lúc demo, nếu GPS hoặc mạng bản đồ không ổn định, người dùng vẫn có thể mở tour bằng QR và chọn quán từ danh sách.")
    add_heading(doc, "4.1. Map rendering và POI detail", 2)
    add_bullets(doc, [
        "Dữ liệu POI được lấy từ API khi online; nếu API lỗi, app dùng IndexedDB hoặc demo fallback để tránh màn hình trắng.",
        "Marker theo category giúp phân biệt hải sản, lẩu, street food, snack, landmark và các điểm dừng trong tour.",
        "showPoiDetail() là điểm hội tụ của nhiều thao tác: click marker, chọn trong tour list, quét QR, next/previous trong tour.",
        "Nội dung POI detail gồm tên quán, địa chỉ, mô tả, giờ mở cửa, mức giá, nút nghe audio và trạng thái dịch theo ngôn ngữ đang chọn.",
    ])
    add_heading(doc, "4.2. GPS/geofence trigger", 2)
    add_para(doc, "Geofence dùng vị trí hiện tại để gợi ý điểm gần người dùng. Vì GPS trong hẻm/quán đông người có thể lệch, thiết kế không khóa trải nghiệm vào GPS. QR Tour tại cổng và danh sách tour là đường fallback chính, giúp demo vẫn ổn khi trình duyệt chặn location hoặc điện thoại chưa cấp quyền.")
    add_table(doc, ["Tình huống", "Cách xử lý", "Ý nghĩa khi demo"], [
        ["GPS hoạt động", "Tính khoảng cách tới POI và gợi ý điểm gần nhất.", "Có thể nói hệ thống hỗ trợ tương tác theo vị trí thực tế."],
        ["GPS bị chặn", "Vẫn mở tour bằng QR hoặc chọn POI từ danh sách.", "Demo không bị phụ thuộc vào quyền location."],
        ["API lỗi", "offlineDB.loadPois() hoặc demo data giữ dữ liệu cơ bản.", "Không bị trắng màn hình trước giảng viên."],
        ["Map tile tải chậm", "PWA cache và danh sách POI vẫn cho phép thao tác.", "Bản đồ là lớp hỗ trợ, không phải điểm lỗi duy nhất."],
    ], [3.8, 6.2, 6.5], 8.8)
    add_heading(doc, "4.3. Truy vết thao tác bản đồ", 2)
    add_table(doc, ["Màn hình", "User thao tác", "Hàm/frontend", "Dữ liệu liên quan"], [
        ["Bản đồ du khách", "Mở trang chủ", "initMap(), renderPoiMarkers()", "POI latitude/longitude/category."],
        ["Danh sách tour", "Chọn quán trong tour", "renderTourPoiList() -> showPoiDetail()", "activeTour.pois theo thứ tự tour."],
        ["POI detail", "Bấm nghe", "getPoiScript() -> AudioManager.playDirect()", "ttsScript, translations, selectedLanguage."],
        ["Fallback offline", "Mất mạng rồi mở lại app", "offlineDB.loadPois()", "IndexedDB/cache từ lần tải trước."],
    ], [3.3, 4.0, 5.0, 4.2], 8.6)
def chapter_6(doc):
    add_heading(doc, "CHƯƠNG 6: ADMIN AUTHENTICATION VÀ QUẢN TRỊ", 1)
    add_para(doc, "CMS Admin gồm form đăng nhập, dashboard và các màn quản trị POI/Tour. Mục tiêu bảo mật của đồ án là demo-level: mật khẩu admin không lưu plaintext; hệ thống băm mật khẩu đầu vào rồi so sánh với passwordHash đã seed/lưu. Báo cáo không tuyên bố có refresh token, rate limit, CSRF hoặc phân quyền nhiều cấp.")
    add_heading(doc, "6.1. Login flow", 2)
    add_table(doc, ["Bước", "Thành phần", "Hàm/API", "Mô tả"], [
        ["1", "admin.html", "submit form", "Admin nhập username/password."],
        ["2", "admin.js", "loginAdmin()", "Gửi POST /api/auth/login và xử lý thông báo lỗi."],
        ["3", "Backend", "AuthController.Login() hoặc demo route Program.cs", "Nhận LoginRequest, chuyển sang logic xác thực tương ứng mode."],
        ["4", "AuthService", "LoginAsync(), VerifyPassword()", "Tính hash password đầu vào và so sánh với passwordHash."],
        ["5", "AdminTokenHelper", "token response", "Trả token demo để CMS lưu sessionStorage và gọi API admin."],
    ], [1.0, 3.2, 5.0, 7.3], 8.6)
    note_box(doc, "Ranh giới bảo mật", "Đây là mức phù hợp demo đồ án: không lưu plaintext và có token admin đơn giản. Nếu triển khai production cần bổ sung refresh token, expiry nghiêm ngặt, rate limit, CSRF/CORS chặt hơn, audit log và phân quyền chi tiết.")
    add_heading(doc, "6.2. CRUD POI", 2)
    add_table(doc, ["Chức năng", "Frontend", "Endpoint", "Backend/data", "Validation/fallback cần chú ý"], [
        ["Danh sách POI", "loadPois()", "GET /api/poi/all hoặc /api/poi", "PoiController/PoiService hoặc demoPois", "Nếu API lỗi cần hiển thị lỗi mềm, không làm vỡ CMS."],
        ["Tạo/sửa POI", "savePoi()", "POST/PUT /api/poi", "MongoDB pois hoặc list demoPois", "Tên, vị trí, category, nội dung thuyết minh cần đủ để app user hiển thị."],
        ["Xóa POI", "deletePoi()", "DELETE /api/poi/{id}", "PoiService hoặc demo route", "Không nên xóa POI đang nằm trong tour nếu demo phụ thuộc tour đó."],
        ["Xem QR POI", "viewQr()", "Dữ liệu QR từ frontend", "poi.id/qrCode", "QR POI là luồng phụ, QR Tour là luồng demo chính."],
    ], [2.6, 3.0, 3.5, 4.3, 3.1], 8.0)
    add_heading(doc, "6.3. CRUD Tour", 2)
    add_para(doc, "Tour là thực thể nối nhiều POI theo thứ tự. Khi tạo hoặc sửa tour, admin chọn danh sách POI; backend phải giữ thứ tự poiIds để QR Tour tại cổng mở đúng tuyến. Đây là khác biệt chính giữa QR Tour và QR POI đơn lẻ.")
    add_table(doc, ["Chức năng", "Frontend", "Endpoint", "Backend/service", "Dữ liệu"], [
        ["Load tour", "loadTours()", "GET /api/tour", "TourController hoặc Program.cs demo", "tours/demoTours."],
        ["Mở form tour", "openTourModal()", "Không gọi API ngay", "Dùng dữ liệu POI đã load", "Danh sách POI để chọn thứ tự."],
        ["Lưu tour", "saveTour()", "POST/PUT /api/tour", "TourService", "tour.poiIds, qrCode, isActive."],
        ["Xem QR tour", "viewTourQr()", "QR URL/code", "Dữ liệu tour", "Mã demo chính: VK-TOUR-GATE-001."],
    ], [2.8, 3.2, 3.6, 4.0, 3.1], 8.2)


def chapter_9(doc):
    add_heading(doc, "CHƯƠNG 9: OFFLINE/PWA VÀ DỮ LIỆU DEMO", 1)
    add_para(doc, "App được thiết kế như PWA để phục vụ bối cảnh demo bằng điện thoại: mạng WiFi có thể yếu, CDN có thể tải chậm, MongoDB Atlas có thể chưa kết nối. Vì vậy hệ thống có nhiều lớp dự phòng: service worker cache app shell, IndexedDB lưu POI và demo in-memory mode trong backend.")
    add_heading(doc, "9.1. Service Worker", 2)
    add_table(doc, ["Nhóm tài nguyên", "Chiến lược", "Ví dụ", "Lý do"], [
        ["App shell", "CacheFirst", "index.html, admin.html, CSS, JS, manifest.json", "Mở lại nhanh, giảm phụ thuộc mạng."],
        ["API GET", "NetworkFirst", "/api/poi, /api/tour/qr/{qrCode}", "Ưu tiên dữ liệu mới, có cache khi lỗi mạng."],
        ["Map tiles", "CacheFirst", "Tile OpenStreetMap", "Giảm tải khi pan/zoom khu vực đã xem."],
        ["CDN", "CacheFirst", "Leaflet, html5-qrcode, QRCode", "Các thư viện demo không phải tải lại liên tục."],
    ], [3.0, 3.0, 5.0, 5.5], 8.4)
    add_heading(doc, "9.2. IndexedDB/cache nội bộ", 2)
    add_bullets(doc, [
        "offlineDB.savePois() lưu danh sách POI sau khi tải thành công từ API.",
        "offlineDB.loadPois() đọc lại POI khi API lỗi hoặc app đang offline.",
        "Translation cache giảm số lần gọi dịch khi người dùng đổi qua lại nhiều ngôn ngữ.",
        "PWA cache không thay thế database; nó chỉ giúp app demo ổn định hơn ở phía client.",
    ])
    add_heading(doc, "9.3. Demo seed data", 2)
    add_table(doc, ["Dữ liệu", "Hàm seed", "Vai trò trong demo"], [
        ["POI/quán ăn", "CreateDemoPois()", "Có sẵn danh sách quán, vị trí, mô tả và script thuyết minh."],
        ["QR Tour cổng", "CreateDemoTours()", "Tạo tour chính với mã VK-TOUR-GATE-001."],
        ["Analytics mẫu", "CreateDemoAnalyticsEvents()", "Dashboard không trắng trước khi có event thật."],
        ["Admin mặc định", "AuthService.DefaultAdminPasswordHash", "Seed bằng hash, không lưu plaintext."],
    ], [3.3, 4.4, 8.8], 8.6)
    note_box(doc, "Khi nào dùng demo mode", "Demo mode dùng khi không có MongoDB connection string hợp lệ. Frontend vẫn gọi cùng endpoint nên cách trình bày với giảng viên không thay đổi; chỉ khác nguồn dữ liệu phía sau là in-memory.")


def chapter_11(doc):
    add_heading(doc, "CHƯƠNG 11: BẢNG ĐẶC TẢ HỆ SINH THÁI CHỨC NĂNG", 1)
    add_para(doc, "Phần này là bảng truy vết chính để trả lời câu hỏi “chức năng này gọi hàm nào”. Để giữ A4 portrait dễ đọc, bảng được tách thành hai phần: truy vết thao tác người dùng và truy vết backend/dữ liệu.")
    add_heading(doc, "11.1. Truy vết thao tác UI -> API", 2)
    add_table(doc, ["Chức năng", "User thao tác", "File frontend", "Hàm frontend", "API endpoint", "Demo nhanh"], [
        ["Admin login", "Nhập tài khoản admin", "wwwroot/js/admin.js", "loginAdmin()", "POST /api/auth/login", "Mở admin.html, đăng nhập, xem token/session CMS."],
        ["CRUD POI", "Thêm/sửa/xóa quán", "wwwroot/js/admin.js", "savePoi(), deletePoi()", "POST/PUT/DELETE /api/poi", "Sửa tên/mô tả POI rồi reload app user."],
        ["CRUD Tour", "Tạo/sửa tour", "wwwroot/js/admin.js", "loadTours(), openTourModal(), saveTour()", "GET/POST/PUT/DELETE /api/tour", "Chọn POI theo thứ tự và xem QR tour."],
        ["QR Tour", "Quét QR tại cổng", "wwwroot/js/app.js", "handleQrCode() -> handleTourQrCode()", "GET /api/tour/qr/{qrCode}", "Quét VK-TOUR-GATE-001, danh sách quán mở theo thứ tự."],
        ["Audio/TTS", "Bấm nghe thuyết minh", "wwwroot/js/app.js, audio-manager.js", "getPoiScript() -> AudioManager.playDirect()", "GET /api/tts khi fallback", "Tắt voice phù hợp hoặc đổi ngôn ngữ để xem fallback."],
        ["Đa ngôn ngữ", "Đổi dropdown ngôn ngữ", "wwwroot/js/app.js", "changeLanguage() -> translateWithCache()", "Dịch runtime/cache", "Đổi sang English/Japanese, UI và nội dung đổi theo."],
    ], [2.3, 2.7, 3.3, 3.2, 3.2, 3.8], 7.5)
    add_heading(doc, "11.2. Truy vết analytics, online và offline", 2)
    add_table(doc, ["Chức năng", "File frontend", "Hàm frontend", "Endpoint", "Backend/service", "Data source"], [
        ["Dashboard stats", "wwwroot/js/admin.js", "loadDashboardData()", "/api/analytics/stats", "AnalyticsService.GetEventCountsAsync()", "analytics/demo analyticsEvents."],
        ["Top POIs", "wwwroot/js/admin.js", "renderTopPoisChart()", "/api/analytics/top-pois", "AnalyticsService.GetTopPoiStatsAsync()", "eventType = poi_listen."],
        ["Recent events", "wwwroot/js/admin.js", "renderRecentEvents()", "/api/analytics/recent", "AnalyticsService.GetRecentEventsAsync()", "analytics theo thời gian."],
        ["Heatmap", "wwwroot/js/admin.js", "initHeatmap()", "/api/analytics/heatmap", "AnalyticsService.GetHeatmapDataAsync()", "location_update hoặc event có lat/lng."],
        ["Presence heartbeat", "wwwroot/js/app.js", "startPresenceHeartbeat() -> sendPresenceHeartbeat()", "POST /api/presence/heartbeat", "UserPresenceService.Upsert()", "in-memory sessions."],
        ["Admin kick", "wwwroot/js/admin.js", "loadOnlineUsers(), kickOnlineUser()", "GET/POST /api/admin/online-users", "UserPresenceService.GetOnlineSessions(), Kick()", "session state in-memory."],
        ["Client kicked", "wwwroot/js/app.js", "handleSessionKicked()", "heartbeat trả kicked = true", "UserPresenceService giữ IsKicked", "sessionId hiện tại."],
        ["Offline/PWA", "wwwroot/js/app.js, sw.js", "offlineDB.savePois(), offlineDB.loadPois()", "API GET khi online", "Service worker + IndexedDB", "cache app shell/POI."],
        ["AAC", "wwwroot/js/app.js", "aacSpeak(), detectLanguage()", "GET /api/tts khi fallback", "AudioManager/TTS route", "câu mẫu frontend."],
    ], [2.4, 3.0, 3.7, 3.2, 4.2, 3.0], 7.3)
    add_heading(doc, "11.3. Chuỗi hàm bắt buộc cần nhớ", 2)
    add_bullets(doc, [
        "QR Tour: QRScannerManager._onScanSuccess() -> handleQrCode() -> handleTourQrCode() -> extractTourQrCode() -> resolveTourByQrCode() -> GET /api/tour/qr/{qrCode} -> openTourFromQr() -> renderTourPoiList() -> showPoiDetail() -> trackTourQrScan().",
        "Dashboard: loadDashboardData() -> /api/analytics/stats, /api/analytics/top-pois, /api/analytics/recent, /api/analytics/heatmap -> renderTopPoisChart() -> renderRecentEvents() -> initHeatmap().",
        "Audio/ngôn ngữ: changeLanguage() -> normalizeAppLanguage() -> applyUILanguage() -> translateWithCache() -> translateText() -> queuePoiTranslationWarmup() -> getPoiScript() -> AudioManager.playDirect() -> _speak() -> _speakWithWebSpeech() hoặc _speakWithGoogleTTS() -> /api/tts.",
        "Online/kick: startPresenceHeartbeat() -> sendPresenceHeartbeat() -> POST /api/presence/heartbeat -> UserPresenceService.Upsert(); admin dùng loadOnlineUsers() và kickOnlineUser(); client nhận handleSessionKicked().",
        "Auth: form login -> loginAdmin() -> POST /api/auth/login -> AuthController hoặc Program.cs demo endpoint -> AuthService -> VerifyPassword() -> token response.",
    ])


def appendix(doc):
    add_heading(doc, "PHỤ LỤC", 1)
    add_heading(doc, "A. Technology stack", 2)
    add_table(doc, ["Nhóm", "Công nghệ", "Ghi chú"], [
        ["Backend", "ASP.NET Core 10, C#, Minimal API, MVC Controller, OpenAPI", "Cùng phục vụ MongoDB mode và demo mode."],
        ["Database", "MongoDB Atlas; demo in-memory", "Không ghi connection string hoặc secret vào báo cáo."],
        ["Frontend", "HTML/CSS/Vanilla JavaScript, PWA, Service Worker, IndexedDB", "Dễ giải thích trong môn C# vì frontend chỉ gọi REST API."],
        ["Map/QR", "Leaflet/OpenStreetMap, html5-qrcode, qrcode.js", "Phù hợp demo miễn phí."],
        ["Audio", "Web Speech API, Google TTS fallback qua /api/tts", "Fallback để tránh crash khi browser thiếu voice."],
        ["Deploy/demo", "dotnet run, localhost/LAN", "Điện thoại cần cùng mạng LAN khi demo thực tế."],
    ], [3.2, 7.5, 5.8], 8.5)
    add_heading(doc, "B. Cost / free-tier analysis", 2)
    add_para(doc, "Bản demo có thể vận hành với chi phí 0 đồng nếu dùng MongoDB Atlas free tier, OpenStreetMap/Leaflet, Web Speech API và cache cục bộ. Tuy nhiên free tier không đồng nghĩa SLA production. Endpoint TTS fallback phụ thuộc mạng và dịch vụ bên ngoài; nếu triển khai thật cần thay bằng dịch vụ có quota, chính sách sử dụng rõ ràng và giám sát lỗi.")
    add_heading(doc, "C. Testing commands", 2)
    add_table(doc, ["Lệnh", "Mục tiêu", "Kỳ vọng"], [
        ["node --check CShape/VinhKhanhFoodTour.Api/wwwroot/js/app.js", "Kiểm tra cú pháp JavaScript app du khách.", "Không có SyntaxError."],
        ["node --check CShape/VinhKhanhFoodTour.Api/wwwroot/js/admin.js", "Kiểm tra cú pháp JavaScript CMS admin.", "Không có SyntaxError."],
        ["dotnet build CShape/VinhKhanhFoodTour.Api/VinhKhanhFoodTour.Api.csproj", "Biên dịch backend ASP.NET Core.", "Build thành công; nếu .exe bị khóa thì dừng process đang chạy."],
        ["dotnet run --project CShape/VinhKhanhFoodTour.Api/VinhKhanhFoodTour.Api.csproj", "Chạy server local để demo.", "Mở được http://127.0.0.1:5000 hoặc URL cấu hình."],
    ], [8.3, 4.7, 3.5], 8.4)
    note_box(doc, "Lỗi build thường gặp", "Nếu dotnet build báo file .exe bị khóa bởi VinhKhanhFoodTour.Api (PID), nghĩa là server đang chạy. Dừng terminal dotnet run hoặc Stop-Process theo PID rồi build lại.")
    add_heading(doc, "D. Demo talking points", 2)
    add_bullets(doc, [
        "Số QR Scans lấy từ eventCounts.qr_scan trong API /api/analytics/stats.",
        "Top POI nghe nhiều nhất là ListenCount, được group từ event poi_listen.",
        "Quét QR gọi handleQrCode() -> handleTourQrCode() -> resolveTourByQrCode() -> GET /api/tour/qr/{qrCode}.",
        "Đổi ngôn ngữ gọi changeLanguage() -> normalizeAppLanguage() -> applyUILanguage() -> translateWithCache().",
        "Kick user: admin gọi API kick; heartbeat sau đó trả kicked = true; client chạy handleSessionKicked().",
        "Password admin không lưu plaintext; hệ thống băm password đầu vào rồi so sánh hash.",
        "Web không có giới hạn cứng số user online trong code; bản demo bị giới hạn thực tế bởi RAM/CPU/network và dữ liệu presence đang nằm in-memory.",
    ])
    add_heading(doc, "E. Placeholder cần điền trước khi nộp", 2)
    add_bullets(doc, [
        "Thành viên nhóm: [Điền tên thành viên].",
        "Mã số sinh viên: [Điền mã số sinh viên].",
        "Giảng viên phụ trách: [Điền tên giảng viên].",
    ])


def add_heading(doc, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    if level == 1 and (text.startswith("CHƯƠNG") or text.startswith("PHỤ LỤC")):
        paragraph.paragraph_format.page_break_before = True
    return paragraph


def add_toc(doc: Document) -> None:
    add_heading(doc, "MỤC LỤC TÓM TẮT", 1)
    chapters = [
        "CHƯƠNG 1: Kiến trúc tổng quan",
        "CHƯƠNG 2: Luồng khởi động hệ thống",
        "CHƯƠNG 3: QR Tour và tương tác O2O",
        "CHƯƠNG 4: Bản đồ, GPS và Geofence",
        "CHƯƠNG 5: Audio, TTS và đa ngôn ngữ",
        "CHƯƠNG 6: Admin Authentication và quản trị",
        "CHƯƠNG 7: Dashboard Analytics",
        "CHƯƠNG 8: Online Users và Kick Session",
        "CHƯƠNG 9: Offline/PWA và dữ liệu demo",
        "CHƯƠNG 10: End-to-End Flow",
        "CHƯƠNG 11: Bảng đặc tả hệ sinh thái chức năng",
        "CHƯƠNG 12: Design Patterns & Key Constants",
        "PHỤ LỤC: Công nghệ, chi phí, kiểm thử và gợi ý bảo vệ",
    ]
    for item in chapters:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def build() -> None:
    figs = create_diagrams()
    doc = Document()
    setup_styles(doc)
    set_header_footer(doc)
    add_cover(doc)
    add_toc(doc)
    chapter_1(doc, figs)
    chapter_2(doc, figs)
    chapter_3(doc, figs)
    chapter_4(doc)
    chapter_5(doc, figs)
    chapter_6(doc)
    chapter_7(doc, figs)
    chapter_8(doc, figs)
    chapter_9(doc)
    chapter_10(doc, figs)
    chapter_11(doc)
    chapter_12(doc)
    appendix(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
